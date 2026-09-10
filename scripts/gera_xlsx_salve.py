# -*- coding: utf-8 -*-
"""
Gera XLSX normalizado (modelo estrela) a partir do CSV consolidado do SALVE.
Uso: python gera_xlsx_salve.py salve_fauna_consolidado.csv salve_fauna.xlsx
"""
import sys, re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ENTRADA = sys.argv[1] if len(sys.argv) > 1 else "salve_fauna_consolidado.csv"
SAIDA = sys.argv[2] if len(sys.argv) > 2 else "salve_fauna.xlsx"
SEP = "; "

df = pd.read_csv(ENTRADA, dtype=str, keep_default_na=False)
# remove caracteres de controle invisíveis (ex.: \x01 dentro de nomes de autor)
_ctrl = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
n_ctrl = int(df.apply(lambda c: c.str.contains(_ctrl)).sum().sum())
df = df.apply(lambda c: c.str.replace(_ctrl, "", regex=True))
print("células com caracteres de controle removidos:", n_ctrl)
for c in ["id_ficha", "ameacada", "n_biomas", "n_estados", "ano_avaliacao"] + \
         [c for c in df.columns if c.startswith(("bioma_", "reg_"))]:
    df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
for c in ["altitude_min", "altitude_max", "batimetria_min", "batimetria_max"]:
    df[c] = pd.to_numeric(df[c], errors="coerce")

ID = ["id_ficha", "nome_cientifico"]


def explode(col, nome_col):
    s = df[ID + [col]].copy()
    s[col] = s[col].str.split(SEP)
    s = s.explode(col)
    s = s[s[col].notna() & (s[col].str.strip() != "")]
    return s.rename(columns={col: nome_col}).reset_index(drop=True)


def explode_codigos(col):
    """'2.3.4 - Pecuária em escala desconhecida' -> codigo, nivel, descricao"""
    s = explode(col, "item")
    m = s["item"].str.extract(r"^\s*([\d.]+)\s*-\s*(.*)$")
    s["codigo"] = m[0].str.strip()
    s["descricao"] = m[1].fillna(s["item"]).str.strip()
    s["nivel"] = s["codigo"].fillna("").str.count(r"\d+")
    s["codigo_pai"] = s["codigo"].str.rsplit(".", n=1).str[0].where(s["nivel"] > 1)
    return s[ID + ["codigo", "nivel", "codigo_pai", "descricao"]]


# ------------------------------------------------------------ abas
fichas = df[[
    "id_ficha", "nome_cientifico", "filo", "classe", "ordem", "familia", "genero", "especie",
    "subespecie", "nome_cientifico_anterior", "autor", "nome_comum", "grupo",
    "ano_avaliacao", "data_avaliacao", "categoria", "categoria_sigla", "ameacada",
    "possivemente_extinta", "criterio", "endemica_brasil", "consta_em_lista_nacional_oficial",
    "migratoria", "tendencia_populacional", "n_biomas", "n_estados",
    "altitude_min", "altitude_max", "batimetria_min", "batimetria_max",
] + [c for c in df.columns if c.startswith("bioma_")] + [c for c in df.columns if c.startswith("reg_")]]

ucs = pd.concat([
    explode("unidade_de_conservacao_federal", "unidade_conservacao").assign(esfera="Federal"),
    explode("unidade_de_conservacao_estadual", "unidade_conservacao").assign(esfera="Estadual"),
    explode("rppn", "unidade_conservacao").assign(esfera="RPPN"),
], ignore_index=True)[ID + ["esfera", "unidade_conservacao"]]

conservacao = pd.concat([
    explode("acao_conservacao", "item").assign(tipo="Ação de conservação"),
    explode("plano_de_acao", "item").assign(tipo="Plano de ação"),
    explode("listas_e_convencoes", "item").assign(tipo="Lista / convenção"),
], ignore_index=True)[ID + ["tipo", "item"]]

abas = {
    "fichas": fichas,
    "estados": explode("estado", "estado"),
    "biomas": explode("bioma", "bioma"),
    "bacias": explode("bacia_hidrografica", "bacia_hidrografica"),
    "ucs": ucs,
    "ameacas": explode_codigos("ameaca"),
    "usos": explode_codigos("uso"),
    "conservacao": conservacao,
    "justificativas": df[ID + ["justificativa"]][df["justificativa"] != ""],
}

dicionario = pd.DataFrame([
    ("fichas", "uma linha por espécie; chave id_ficha; dummies bioma_*/reg_* (0/1) prontas para filtro e soma"),
    ("estados", "relação N:N espécie x estado (relate por id_ficha)"),
    ("biomas", "relação N:N espécie x bioma"),
    ("bacias", "relação N:N espécie x sub-bacia hidrográfica"),
    ("ucs", "relação N:N espécie x unidade de conservação; esfera = Federal / Estadual / RPPN"),
    ("ameacas", "ameaças classificadas pelo código IUCN (nivel 1 = categoria, 2 = subcategoria, 3+ = detalhe); codigo_pai liga a hierarquia"),
    ("usos", "usos da espécie, mesma estrutura hierárquica de ameacas"),
    ("conservacao", "ações de conservação, planos de ação (PAN) e listas/convenções em que a espécie consta"),
    ("justificativas", "texto integral da justificativa da avaliação"),
    ("", ""),
    ("categoria_sigla", "EX Extinta, EW Extinta na Natureza, RE Regionalmente Extinta, CR Criticamente em Perigo, EN Em Perigo, VU Vulnerável, NT Quase Ameaçada, LC Menos Preocupante, DD Dados Insuficientes, NA Não Aplicável"),
    ("ameacada", "1 quando categoria em CR/EN/VU"),
    ("fonte", "SALVE / ICMBio, exportação pública das fichas por bioma em 09/09/2026; 26.818 linhas brutas deduplicadas para 15.305 espécies"),
], columns=["item", "descricao"])
abas["dicionario"] = dicionario

# ------------------------------------------------------------ gravação
with pd.ExcelWriter(SAIDA, engine="openpyxl") as w:
    for nome, t in abas.items():
        t.to_excel(w, sheet_name=nome, index=False)
        print(f"{nome:15} {len(t):6} linhas x {t.shape[1]} colunas")

# ------------------------------------------------------------ formatação
wb = load_workbook(SAIDA)
hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E78")
body = Font(name="Arial", size=10)
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        larg = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(larg + 2, 10), 60)
wb.save(SAIDA)
print("gravado", SAIDA)
