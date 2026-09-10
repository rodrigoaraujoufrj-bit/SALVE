# -*- coding: utf-8 -*-
"""
Cruza os anexos das Portarias MMA 1.667/2026 (peixes e invertebrados aquáticos) e
1.704/2026 (demais grupos) com o consolidado do SALVE e grava no XLSX, com realce:
  amarelo  = espécie nova na lista (asterisco na portaria)
  laranja  = categoria da portaria diferente da categoria no SALVE
  vermelho = ameaçada no SALVE mas ausente das duas portarias
Uso: python cruza_portarias.py salve_fauna_consolidado.csv salve_fauna.xlsx
"""
import sys, re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CSV, XLSX = sys.argv[1], sys.argv[2]
PORTARIAS = {
    "1667": {"arquivo": "dados/portarias/portaria_1667_anexo1.txt", "lista_anterior": "Portaria MMA 445/2014",
             "grupos": {"PC": "Peixes Continentais", "PM": "Peixes Marinhos (ósseos)", "TR": "Tubarões e Raias",
                        "ID": "Invertebrados de Água Doce", "IM": "Invertebrados Marinhos"}},
    "1704": {"arquivo": "dados/portarias/portaria_1704_anexos.txt", "lista_anterior": "Portaria MMA 148/2022",
             "grupos": {"AN": "Anfíbios", "AV": "Aves", "IT": "Invertebrados Terrestres",
                        "MA": "Mamíferos", "RE": "Répteis"}},
}
MANUAL = {"Sartor tucuruiensis": "Sartor tucuruiense", "Girardia paucipunctata": "Girardia paucipuntacta"}
AMARELO, LARANJA, VERMELHO = "FFF2A8", "F8CBAD", "F4B6B6"
norm = lambda x: re.sub(r"\s+", "", str(x)).lower()

s = pd.read_csv(CSV, dtype=str, keep_default_na=False)
s["id_ficha"] = s["id_ficha"].astype(int)
s["k"] = s["nome_cientifico"].map(norm)
s["k_esp"] = s["especie"].map(norm)
s["epi"] = s["nome_cientifico"].str.split().str[-1].str.lower()


def casar(r):
    k = norm(r.especie_portaria)
    c = s[s.k == k]
    if len(c) == 1: return c.iloc[0], "nome idêntico"
    c = s[s.k_esp == k]                      # portaria sem subespécie, SALVE com
    if len(c) == 1: return c.iloc[0], "nome idêntico (subespécie no SALVE)"
    if r.especie_portaria in MANUAL:
        c = s[s.especie == MANUAL[r.especie_portaria]]
        if len(c) == 1: return c.iloc[0], "grafia divergente"
    partes = r.especie_portaria.split()
    c = s[(s.epi == partes[-1].lower()) & (s.familia == r.familia_portaria)]
    if len(partes) == 3:                      # trinômio: exige mesmo gênero+espécie
        c = c[c.especie.str.lower() == " ".join(partes[:2]).lower()]
    if len(c) == 1: return c.iloc[0], "gênero revisado (sinônimo)"
    return None, "não consta nas exportações por bioma"


abas, todos = {}, []
for num, cfg in PORTARIAS.items():
    p = pd.read_csv(cfg["arquivo"], sep="|", dtype=str, keep_default_na=False)
    p = p.rename(columns={"n": "n_portaria", "familia": "familia_portaria",
                          "especie": "especie_portaria", "cat": "categoria_portaria"})
    p["anexo"] = p.n_portaria.str.match(r"^\d+$").map({True: "I", False: "II (extintas)"})
    p["novo_na_lista"] = (p.pop("novo") == "*").astype(int)
    p["grupo_portaria"] = p.pop("grupo").map(cfg["grupos"])
    p["categoria_sigla_portaria"] = p.categoria_portaria.str.replace(r"\s*\(PE\)", "", regex=True)
    p["possivelmente_extinta_portaria"] = p.categoria_portaria.str.contains("PE").astype(int)
    res = []
    for _, r in p.iterrows():
        c, tipo = casar(r)
        res.append((None, "", "", "", tipo) if c is None else
                   (c.id_ficha, c.nome_cientifico, c.categoria_sigla, c.consta_em_lista_nacional_oficial, tipo))
    p[["id_ficha", "especie_salve", "categoria_salve", "consta_lista_salve", "tipo_match"]] = pd.DataFrame(res, index=p.index)
    p["id_ficha"] = p.id_ficha.astype("Int64")
    p["categoria_divergente"] = ((p.categoria_salve != "") & (p.categoria_salve != p.categoria_sigla_portaria)).astype(int)
    p["portaria"] = num
    p = p[["portaria", "anexo", "n_portaria", "novo_na_lista", "grupo_portaria", "familia_portaria", "especie_portaria",
           "categoria_portaria", "categoria_sigla_portaria", "possivelmente_extinta_portaria",
           "id_ficha", "especie_salve", "categoria_salve", "categoria_divergente", "consta_lista_salve", "tipo_match"]]
    abas["portaria_" + num] = p; todos.append(p)
    print(f"portaria {num}: {len(p)} táxons | novos {p.novo_na_lista.sum()} | divergentes {p.categoria_divergente.sum()} | sem match {(p.id_ficha.isna()).sum()}")

todos = pd.concat(todos, ignore_index=True)
lk = todos[todos.id_ficha.notna()].drop_duplicates("id_ficha").set_index("id_ficha")

# ameaçadas no SALVE que não aparecem em nenhuma portaria
fora = s[(s.ameacada == "1") & ~s.id_ficha.isin(lk.index)]
fora = fora[["id_ficha", "nome_cientifico", "grupo", "classe", "familia", "categoria_sigla", "ano_avaliacao",
             "consta_em_lista_nacional_oficial", "bioma"]].rename(columns={"categoria_sigla": "categoria_salve"})
fora["situacao"] = fora.consta_em_lista_nacional_oficial.map({
    "Sim": "constava na lista anterior e saiu das portarias 2026",
    "Não": "avaliada como ameaçada no SALVE mas ainda não oficializada"}).fillna("ausente das portarias 2026")
print("ameaçadas no SALVE fora das portarias:", len(fora), fora.grupo.value_counts().to_dict())

# resumo das mudanças
mud = pd.concat([
    todos[todos.novo_na_lista == 1].assign(tipo_mudanca="Nova na lista"),
    todos[todos.categoria_divergente == 1].assign(tipo_mudanca="Categoria diferente do SALVE"),
], ignore_index=True)[["tipo_mudanca", "portaria", "anexo", "n_portaria", "grupo_portaria", "especie_portaria",
                       "categoria_portaria", "categoria_salve", "especie_salve", "id_ficha", "tipo_match"]]
mud = pd.concat([mud, fora.rename(columns={"nome_cientifico": "especie_salve", "grupo": "grupo_portaria"})
                 .assign(tipo_mudanca="Ameaçada no SALVE, fora das portarias")[
                     ["tipo_mudanca", "grupo_portaria", "especie_salve", "categoria_salve", "id_ficha"]]], ignore_index=True)
abas["mudancas"] = mud
abas["salve_fora_portarias"] = fora

with pd.ExcelWriter(XLSX, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
    fichas = pd.read_excel(XLSX, sheet_name="fichas")
    fichas = fichas.drop(columns=[c for c in fichas.columns if "portaria" in c], errors="ignore")
    fichas["portaria"] = fichas.id_ficha.map(lk.portaria).fillna("")
    fichas["categoria_portaria"] = fichas.id_ficha.map(lk.categoria_portaria).fillna("")
    fichas["novo_na_portaria"] = fichas.id_ficha.map(lk.novo_na_lista).fillna(0).astype(int)
    fichas["categoria_divergente_portaria"] = fichas.id_ficha.map(lk.categoria_divergente).fillna(0).astype(int)
    fichas["ameacada_fora_portarias"] = fichas.id_ficha.isin(fora.id_ficha).astype(int)
    fichas.to_excel(w, sheet_name="fichas", index=False)
    for nome, t in abas.items():
        t.to_excel(w, sheet_name=nome, index=False)
    dic = pd.read_excel(XLSX, sheet_name="dicionario")
    dic = dic[~dic.item.astype(str).str.contains("portaria|mudancas|realce", case=False)]
    dic = pd.concat([dic, pd.DataFrame([
        ("portaria_1667", "Anexo I da Portaria GM/MMA 1.667/2026 (peixes e invertebrados aquáticos, 490 táxons, DOU 28/04/2026); asterisco = não constava na Portaria 445/2014"),
        ("portaria_1704", "Anexos I e II da Portaria MMA 1.704/2026 (anfíbios, aves, invertebrados terrestres, mamíferos, répteis; 790 ameaçadas + 9 extintas, DOU 17/06/2026); asterisco = não constava na Portaria 148/2022"),
        ("mudancas", "lista única de tudo que mudou: novas na lista, categoria diferente do SALVE, ameaçadas no SALVE ausentes das portarias"),
        ("salve_fora_portarias", "espécies com categoria CR/EN/VU no SALVE que não aparecem em nenhuma das duas portarias"),
        ("realce", "amarelo = nova na lista; laranja = categoria da portaria difere do SALVE; vermelho = ameaçada no SALVE fora das portarias"),
        ("fichas (portarias)", "colunas portaria / categoria_portaria / novo_na_portaria / categoria_divergente_portaria / ameacada_fora_portarias"),
    ], columns=["item", "descricao"])], ignore_index=True)
    dic.to_excel(w, sheet_name="dicionario", index=False)

# ---------------- formatação e realce
wb = load_workbook(XLSX)
def fmt(ws):
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font = Font(name="Arial", size=10)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        larg = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(larg + 2, 10), 60)

def realcar(ws, col_novo=None, col_div=None, col_fora=None):
    hdr = {c.value: i for i, c in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        cor = None
        if col_fora and row[hdr[col_fora]].value == 1: cor = VERMELHO
        if col_div and row[hdr[col_div]].value == 1: cor = LARANJA
        if col_novo and row[hdr[col_novo]].value == 1: cor = AMARELO
        if col_novo and col_div and row[hdr[col_novo]].value == 1 and row[hdr[col_div]].value == 1: cor = LARANJA
        if cor:
            for cell in row: cell.fill = PatternFill("solid", fgColor=cor)

for nome in ["fichas", "portaria_1667", "portaria_1704", "mudancas", "salve_fora_portarias", "dicionario"]:
    fmt(wb[nome])
realcar(wb["fichas"], "novo_na_portaria", "categoria_divergente_portaria", "ameacada_fora_portarias")
for n in ["portaria_1667", "portaria_1704"]:
    realcar(wb[n], "novo_na_lista", "categoria_divergente")
ws = wb["mudancas"]; hdr = {c.value: i for i, c in enumerate(ws[1])}
cores = {"Nova na lista": AMARELO, "Categoria diferente do SALVE": LARANJA, "Ameaçada no SALVE, fora das portarias": VERMELHO}
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.fill = PatternFill("solid", fgColor=cores[row[hdr["tipo_mudanca"]].value])
for row in wb["salve_fora_portarias"].iter_rows(min_row=2):
    for cell in row: cell.fill = PatternFill("solid", fgColor=VERMELHO)
ordem = ["fichas", "portaria_1667", "portaria_1704", "mudancas", "salve_fora_portarias"]
wb._sheets = [wb[n] for n in ordem] + [ws for ws in wb.worksheets if ws.title not in ordem]
wb.save(XLSX); print("gravado", XLSX)
