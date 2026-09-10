# -*- coding: utf-8 -*-
"""
Monta a tabela final de status por espécie (SALVE x Portarias 2026 x IUCN global) a partir do
consolidado, do XLSX já cruzado com as portarias e do cache iucn_global.csv.
Uso: python scripts/status_comparado.py saida/salve_fauna_consolidado.csv saida/salve_fauna.xlsx saida/iucn_global.csv
"""
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CSV, XLSX, CACHE = sys.argv[1:4]
ORDEM = ["EX", "EW", "RE", "CR", "EN", "VU", "NT", "LC", "DD", "NA", "NE"]
rank = {c: i for i, c in enumerate(ORDEM)}
AMARELO, LARANJA, VERMELHO, AZUL, CINZA = "FFF2A8", "F8CBAD", "F4B6B6", "CFE2F3", "EDEDED"

s = pd.read_csv(CSV, dtype=str, keep_default_na=False)
s = s.drop(columns=[c for c in s.columns if "iucn_global" in c], errors="ignore")
g = pd.read_csv(CACHE, dtype=str, keep_default_na=False)
g.columns = [c.lstrip("\ufeff") for c in g.columns]

# HIGHERRANK = GBIF só achou o gênero: não há categoria válida
hr = g.gbif_match == "HIGHERRANK"
g.loc[hr, ["iucn_codigo", "categoria_iucn_global", "gbif_nome_aceito"]] = ""
g.loc[hr, "gbif_match"] = "sem match (só gênero)"
g.loc[g.iucn_codigo == "NE", "categoria_iucn_global"] = "sem categoria no GBIF"

lk = g.set_index("nome_cientifico")
s["categoria_iucn_global_sigla"] = s.nome_cientifico.map(lk.iucn_codigo).fillna("")
s["categoria_iucn_global"] = s.nome_cientifico.map(lk.categoria_iucn_global).fillna("")
s["iucn_global_nome_aceito"] = s.nome_cientifico.map(lk.gbif_nome_aceito).fillna("")
s["iucn_global_match"] = s.nome_cientifico.map(lk.gbif_match).fillna("")
s.to_csv(CSV, index=False, encoding="utf-8-sig", lineterminator="\n")

# portarias: vêm da aba fichas já cruzada
f = pd.read_excel(XLSX, sheet_name="fichas", dtype={"id_ficha": int})
f = f.drop(columns=[c for c in f.columns if "iucn_global" in c], errors="ignore")
for c in ["categoria_iucn_global_sigla", "iucn_global_nome_aceito"]:
    f[c] = f.id_ficha.map(s.set_index("id_ficha")[c].rename(index=int)).fillna("")
f = f.rename(columns={"categoria_iucn_global_sigla": "categoria_iucn_global"})
pf = f.set_index("id_ficha")

t = s[["id_ficha", "nome_cientifico", "nome_comum", "grupo", "classe", "ordem", "familia", "endemica_brasil",
       "categoria_sigla", "ano_avaliacao", "categoria_iucn_global_sigla", "iucn_global_nome_aceito", "iucn_global_match"]].copy()
t["id_ficha"] = t.id_ficha.astype(int)
t["ano_avaliacao"] = t.ano_avaliacao.apply(lambda v: "" if v in ("", None) else str(int(float(v))))
t = t.rename(columns={"categoria_sigla": "status_salve", "ano_avaliacao": "ano_avaliacao_salve",
                      "categoria_iucn_global_sigla": "status_iucn_global"})
t["portaria_2026"] = t.id_ficha.map(pf.portaria).apply(lambda v: "" if pd.isna(v) or v == "" else str(int(float(v))))
t["status_portaria_2026"] = t.id_ficha.map(pf.categoria_portaria).fillna("")
t.loc[t.status_portaria_2026 == "", "status_portaria_2026"] = "não listada"
t["nova_na_lista_2026"] = t.id_ficha.map(pf.novo_na_portaria).fillna(0).astype(int)
t["ameacada_salve_fora_portarias"] = t.id_ficha.map(pf.ameacada_fora_portarias).fillna(0).astype(int)
sp = t.status_portaria_2026.str.replace(r"\s*\(PE\)", "", regex=True)
t["mudou_salve_x_portaria"] = ((sp != "não listada") & (sp != t.status_salve)).astype(int)

def nac_glob(r):
    n, g = r.status_salve, r.status_iucn_global
    if g in ("", "NE") or n == "": return "sem avaliação global"
    if n == g: return "igual"
    if n in ("DD", "NA") or g in ("DD", "NA"): return "não comparável (DD/NA)"
    return "nacional mais ameaçada" if rank.get(n, 99) < rank.get(g, 99) else "global mais ameaçada"
t["nacional_x_global"] = t.apply(nac_glob, axis=1)
t["mudou_nacional_x_global"] = t.nacional_x_global.isin(["nacional mais ameaçada", "global mais ameaçada"]).astype(int)
t = t[["id_ficha", "nome_cientifico", "nome_comum", "grupo", "classe", "ordem", "familia", "endemica_brasil",
       "status_salve", "ano_avaliacao_salve", "status_portaria_2026", "portaria_2026", "nova_na_lista_2026",
       "status_iucn_global", "iucn_global_nome_aceito", "iucn_global_match",
       "mudou_salve_x_portaria", "ameacada_salve_fora_portarias", "nacional_x_global", "mudou_nacional_x_global"]]
print("status_portaria_2026:", t.status_portaria_2026.value_counts().to_dict())
print("status_iucn_global:", t.status_iucn_global.value_counts().to_dict())
print("nacional_x_global:", t.nacional_x_global.value_counts().to_dict())
t.to_csv(CSV.replace("salve_fauna_consolidado.csv", "status_comparado.csv"), index=False, encoding="utf-8-sig", lineterminator="\n")

with pd.ExcelWriter(XLSX, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
    f.to_excel(w, sheet_name="fichas", index=False)
    t.to_excel(w, sheet_name="status_comparado", index=False)
    dic = pd.read_excel(XLSX, sheet_name="dicionario")
    dic = dic[~dic["item"].astype(str).str.contains("iucn|nacional_x_global|status_comparado", case=False)]
    dic = pd.concat([dic, pd.DataFrame([
        ("status_comparado", "uma linha por espécie com os três status: SALVE (avaliação vigente do ICMBio, com ano), Portaria 2026 (1.667 ou 1.704; 'não listada' quando não consta) e IUCN global (via GBIF)"),
        ("categoria_iucn_global", "categoria da Lista Vermelha global da IUCN obtida via GBIF; NE = sem categoria no GBIF (não avaliada pela IUCN ou nome não casado com o táxon avaliado); vazio = nome não encontrado no GBIF"),
        ("nacional_x_global", "comparação SALVE x IUCN; 'nacional mais ameaçada' ou 'global mais ameaçada' indica qual das duas é mais restritiva"),
    ], columns=["item", "descricao"])], ignore_index=True)
    dic.to_excel(w, sheet_name="dicionario", index=False)

wb = load_workbook(XLSX)
if "nacional_x_global" in wb.sheetnames: del wb["nacional_x_global"]
for nome in ["fichas", "status_comparado", "dicionario"]:
    ws = wb[nome]
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font = Font(name="Arial", size=10)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        larg = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(larg + 2, 10), 60)
# realce: status_portaria em amarelo (nova) / laranja (categoria mudou) / vermelho (fora); status_iucn em azul quando difere
ws = wb["status_comparado"]; h = {c.value: i for i, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    cp = row[h["status_portaria_2026"]]; ci = row[h["status_iucn_global"]]
    if row[h["ameacada_salve_fora_portarias"]].value == 1: cp.fill = PatternFill("solid", fgColor=VERMELHO)
    if row[h["mudou_salve_x_portaria"]].value == 1: cp.fill = PatternFill("solid", fgColor=LARANJA)
    if row[h["nova_na_lista_2026"]].value == 1 and row[h["mudou_salve_x_portaria"]].value == 0: cp.fill = PatternFill("solid", fgColor=AMARELO)
    if row[h["mudou_nacional_x_global"]].value == 1: ci.fill = PatternFill("solid", fgColor=AZUL)
    elif ci.value in ("NE", "", None): ci.fill = PatternFill("solid", fgColor=CINZA)
ordem = ["status_comparado", "fichas", "portaria_1667", "portaria_1704", "mudancas", "salve_fora_portarias"]
wb._sheets = [wb[n] for n in ordem if n in wb.sheetnames] + [w for w in wb.worksheets if w.title not in ordem]
wb.save(XLSX); print("gravado", XLSX)
