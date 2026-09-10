# -*- coding: utf-8 -*-
"""
Busca a categoria global da Lista Vermelha da IUCN para cada espécie do consolidado
do SALVE, via API pública do GBIF (sem token), e grava:
  - iucn_global.csv               : cache com o resultado por espécie (reaproveitado em rodadas seguintes)
  - salve_fauna_consolidado.csv   : ganha as colunas categoria_iucn_global e iucn_global_nome_aceito
  - salve_fauna.xlsx              : aba fichas ganha as mesmas colunas + aba nacional_x_global

Uso:  python iucn_global_gbif.py salve_fauna_consolidado.csv salve_fauna.xlsx
Requer: pip install requests pandas openpyxl
Tempo: ~15 mil espécies, 2 chamadas por espécie, 8 threads -> em torno de 25 a 40 min.
Pode interromper e rodar de novo: o cache evita repetir o que já foi buscado.
"""
import sys, os, time, json, re
from concurrent.futures import ThreadPoolExecutor, as_completed
import pandas as pd
try:                       # rede corporativa com proxy (Netskope): usa o repositório de certificados do Windows
    import truststore; truststore.inject_into_ssl()
except ImportError:
    pass
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CSV = sys.argv[1] if len(sys.argv) > 1 else "salve_fauna_consolidado.csv"
XLSX = sys.argv[2] if len(sys.argv) > 2 else "salve_fauna.xlsx"
CACHE = "iucn_global.csv"
GBIF = "https://api.gbif.org/v1"
THREADS = 8
SESSION = requests.Session()
SESSION.headers["User-Agent"] = "RSAGeo-SALVE-IUCN/1.0"

IUCN_PT = {"EX": "Extinta", "EW": "Extinta na Natureza", "CR": "Criticamente em Perigo", "EN": "Em Perigo",
           "VU": "Vulnerável", "NT": "Quase Ameaçada", "LC": "Menos Preocupante", "DD": "Dados Insuficientes",
           "NE": "Não Avaliada", "NA": "Não Aplicável"}
ORDEM = ["EX", "EW", "CR", "EN", "VU", "NT", "LC", "DD", "NA", "NE"]


def get(url, params=None, tentativas=4):
    for i in range(tentativas):
        try:
            r = SESSION.get(url, params=params, timeout=30)
            if r.status_code == 200:
                return r.json()
            if r.status_code == 404:
                return None
        except requests.RequestException:
            pass
        time.sleep(1.5 * (i + 1))
    return None


def consultar(nome):
    """retorna dict com chave GBIF, nome aceito, tipo de match e categoria IUCN"""
    out = {"nome_cientifico": nome, "gbif_key": "", "gbif_nome_aceito": "", "gbif_match": "",
           "gbif_confianca": "", "categoria_iucn_global": "", "iucn_codigo": "", "erro": ""}
    # nomes provisórios (sp. nov., cf., aff.) não têm avaliação
    if re.search(r"\bsp\.|\bcf\.|\baff\.|'", nome):
        out["gbif_match"] = "nome provisório"; return out
    m = get(f"{GBIF}/species/match", {"name": nome, "kingdom": "Animalia", "strict": "false"})
    if not m or m.get("matchType") in (None, "NONE"):
        out["gbif_match"] = "sem match"; return out
    if m.get("matchType") == "HIGHERRANK":
        # só achou o gênero: tenta a busca textual, que enxerga sinônimos sob outro gênero (ex.: Pauxi mitu -> Mitu mitu)
        sr = get(f"{GBIF}/species/search", {"q": nome, "rank": "SPECIES", "limit": 20})
        cand = [r for r in (sr or {}).get("results", []) if r.get("canonicalName", "").lower() == nome.lower()]
        if not cand:
            out["gbif_match"] = "sem match (só gênero)"; return out
        c = cand[0]
        m = {"matchType": "SEARCH", "confidence": 90, "usageKey": c.get("key"),
             "acceptedUsageKey": c.get("acceptedKey"), "canonicalName": c.get("canonicalName", "")}
    out["gbif_match"] = m.get("matchType", "")
    out["gbif_confianca"] = m.get("confidence", "")
    if m.get("matchType") == "FUZZY" and m.get("confidence", 0) < 90:
        out["gbif_match"] = "fuzzy fraco"; return out
    # se o nome é sinônimo no GBIF, a avaliação fica no táxon aceito
    key = m.get("acceptedUsageKey") or m.get("usageKey")
    out["gbif_key"] = key
    out["gbif_nome_aceito"] = m.get("canonicalName", "")
    if m.get("acceptedUsageKey"):
        acc = get(f"{GBIF}/species/{key}")
        if acc: out["gbif_nome_aceito"] = acc.get("canonicalName", out["gbif_nome_aceito"])
    r = get(f"{GBIF}/species/{key}/iucnRedListCategory")
    if r and r.get("code"):
        out["iucn_codigo"] = r["code"]
        out["categoria_iucn_global"] = IUCN_PT.get(r["code"], r.get("category", r["code"]))
    else:
        out["iucn_codigo"] = "NE"
        out["categoria_iucn_global"] = "Não Avaliada"
    return out


# ------------------------------------------------------------ 1. buscar
df = pd.read_csv(CSV, dtype=str, keep_default_na=False)
nomes = df["nome_cientifico"].drop_duplicates().tolist()
cache = pd.read_csv(CACHE, dtype=str, keep_default_na=False) if os.path.exists(CACHE) else pd.DataFrame(columns=["nome_cientifico"])
prontos = set(cache.loc[~cache.get("gbif_match", pd.Series(dtype=str)).isin(["HIGHERRANK", "sem match", "sem match (só gênero)"]), "nome_cientifico"]) if len(cache) else set()
pendentes = [n for n in nomes if n not in prontos]
cache = cache[cache["nome_cientifico"].isin(prontos)]
print(f"{len(nomes)} espécies, {len(pendentes)} a consultar (cache: {len(cache)})")

novos, t0 = [], time.time()
with ThreadPoolExecutor(THREADS) as ex:
    futs = {ex.submit(consultar, n): n for n in pendentes}
    for i, f in enumerate(as_completed(futs), 1):
        try:
            novos.append(f.result())
        except Exception as e:
            novos.append({"nome_cientifico": futs[f], "erro": str(e)})
        if i % 200 == 0 or i == len(pendentes):
            pd.concat([cache, pd.DataFrame(novos)], ignore_index=True).to_csv(CACHE, index=False, encoding="utf-8-sig")
            print(f"  {i}/{len(pendentes)}  {time.time()-t0:.0f}s")
cache = pd.concat([cache, pd.DataFrame(novos)], ignore_index=True).drop_duplicates("nome_cientifico", keep="last")
cache.to_csv(CACHE, index=False, encoding="utf-8-sig")
print("match GBIF:", cache["gbif_match"].value_counts().to_dict())
print("IUCN global:", cache["iucn_codigo"].value_counts().to_dict())

# ------------------------------------------------------------ 2. anexar ao consolidado
lk = cache.set_index("nome_cientifico")
for c in ["categoria_iucn_global", "iucn_codigo", "gbif_nome_aceito", "gbif_match"]:
    df[c] = df["nome_cientifico"].map(lk[c]).fillna("")
df = df.rename(columns={"iucn_codigo": "categoria_iucn_global_sigla", "gbif_nome_aceito": "iucn_global_nome_aceito",
                        "gbif_match": "iucn_global_match"})
df.to_csv(CSV, index=False, encoding="utf-8-sig", lineterminator="\n")

# ------------------------------------------------------------ 3. XLSX: fichas + aba nacional_x_global
comp = df[["id_ficha", "nome_cientifico", "grupo", "classe", "familia", "ano_avaliacao", "categoria_sigla",
           "categoria_iucn_global_sigla", "iucn_global_nome_aceito", "iucn_global_match", "endemica_brasil"]].copy()
comp = comp.rename(columns={"categoria_sigla": "categoria_nacional_salve", "categoria_iucn_global_sigla": "categoria_iucn_global"})
rank = {c: i for i, c in enumerate(ORDEM)}
def situacao(r):
    n, g = r.categoria_nacional_salve, r.categoria_iucn_global
    if g in ("", "NE") or n == "": return "sem avaliação global"
    if n == g: return "igual"
    if n in ("DD", "NA") or g in ("DD", "NA"): return "não comparável (DD/NA)"
    return "nacional mais ameaçada" if rank.get(n, 99) < rank.get(g, 99) else "global mais ameaçada"
comp["situacao"] = comp.apply(situacao, axis=1)
print("nacional x global:", comp["situacao"].value_counts().to_dict())

with pd.ExcelWriter(XLSX, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
    fichas = pd.read_excel(XLSX, sheet_name="fichas")
    fichas = fichas.drop(columns=[c for c in fichas.columns if "iucn_global" in c], errors="ignore")
    m = df.set_index("id_ficha")
    fichas["id_ficha"] = fichas["id_ficha"].astype(str)
    fichas["categoria_iucn_global"] = fichas["id_ficha"].map(m["categoria_iucn_global_sigla"]).fillna("")
    fichas["iucn_global_nome_aceito"] = fichas["id_ficha"].map(m["iucn_global_nome_aceito"]).fillna("")
    fichas["id_ficha"] = fichas["id_ficha"].astype(int)
    fichas.to_excel(w, sheet_name="fichas", index=False)
    comp.to_excel(w, sheet_name="nacional_x_global", index=False)
    dic = pd.read_excel(XLSX, sheet_name="dicionario")
    dic = dic[~dic["item"].astype(str).str.contains("iucn|nacional_x_global", case=False)]
    dic = pd.concat([dic, pd.DataFrame([
        ("categoria_iucn_global", "categoria da Lista Vermelha global da IUCN, obtida via API do GBIF (data da consulta = data de execução do script); NE = sem avaliação global"),
        ("nacional_x_global", "comparação categoria nacional (SALVE) x global (IUCN); situacao indica qual das duas é mais restritiva"),
    ], columns=["item", "descricao"])], ignore_index=True)
    dic.to_excel(w, sheet_name="dicionario", index=False)

wb = load_workbook(XLSX)
AZUL, CINZA = "CFE2F3", "EDEDED"
for nome in ["fichas", "nacional_x_global", "dicionario"]:
    ws = wb[nome]
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font = Font(name="Arial", size=10)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        larg = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(larg + 2, 10), 60)
ws = wb["nacional_x_global"]; hdr = {c.value: i for i, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    s = row[hdr["situacao"]].value
    cor = AZUL if s in ("nacional mais ameaçada", "global mais ameaçada") else (CINZA if s == "sem avaliação global" else None)
    if cor:
        for cell in row: cell.fill = PatternFill("solid", fgColor=cor)
wb.save(XLSX)
print("gravado", CSV, "e", XLSX)
